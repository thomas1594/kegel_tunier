import io
import os
import sqlite3
import tempfile
import unittest

# Use a temporary database for all tests. Must be set before importing functions.
os.environ["DB_DIR"] = tempfile.mkdtemp(prefix="kegel-test-")

import pandas as pd
from openpyxl import load_workbook

import functions


def make_players(team_scores):
    """Build a player DataFrame from ``{team: [(name, v_total, a_total), ...]}``."""
    rows = []
    player_id = 1
    for team, players in team_scores.items():
        for name, volle, abraeumen in players:
            row = {
                "_id": player_id,
                "Team": team,
                "Spieler": name,
                "Geschlecht": "m",
            }
            player_id += 1
            for i in range(1, 11):
                row[f"V{i}"] = volle // 10 + (1 if i <= volle % 10 else 0)
                row[f"A{i}"] = abraeumen // 10 + (1 if i <= abraeumen % 10 else 0)
            rows.append(row)
    return pd.DataFrame(rows)


class RankingTests(unittest.TestCase):
    def test_calculate_totals(self):
        df = make_players({"A": [("P1", 30, 20)]})
        result = functions.calculate_totals(df)
        self.assertEqual(int(result.loc[0, "Volle"]), 30)
        self.assertEqual(int(result.loc[0, "Abräumen"]), 20)
        self.assertEqual(int(result.loc[0, "Gesamt"]), 50)

    def test_team_tie_break_uses_spare_pins(self):
        df = make_players(
            {
                "Team A": [("P1", 40, 10)],
                "Team B": [("P2", 30, 20)],
            }
        )
        df = functions.calculate_totals(df)
        ranking = functions.get_team_scores(df)
        self.assertEqual(ranking.loc[0, "Team"], "Team B")
        self.assertEqual(ranking.loc[1, "Team"], "Team A")

    def test_player_ranking_limit_and_order(self):
        df = make_players(
            {
                "Team A": [("Low", 10, 5), ("High", 50, 20)],
                "Team B": [("Mid", 30, 10)],
            }
        )
        df = functions.calculate_totals(df)
        ranking = functions.get_player_ranking(df, limit=2)
        self.assertEqual(list(ranking["Spieler"]), ["High", "Mid"])
        self.assertEqual(list(ranking.index), [1, 2])


class TeamResizeTests(unittest.TestCase):
    def test_resize_keeps_renamed_team(self):
        df = make_players({f"Team {i}": [(f"P{i}", 0, 0)] for i in range(1, 11)})
        df.loc[df["Team"] == "Team 1", "Team"] = "Alpha"

        resized = functions.resize_teams(df, 5)

        self.assertEqual(len(functions.sorted_teams(resized)), 5)
        self.assertIn("Alpha", set(resized["Team"]))
        self.assertNotIn("Team 10", set(resized["Team"]))

    def test_resize_adds_unique_team_names(self):
        df = make_players({"Team 1": [("P1", 0, 0)]})
        resized = functions.resize_teams(df, 3)
        self.assertEqual(len(functions.sorted_teams(resized)), 3)
        self.assertEqual(set(resized["Team"]), {"Team 1", "Team 2", "Team 3"})


class KoSyncTests(unittest.TestCase):
    def test_scores_survive_rename_but_reset_on_player_change(self):
        assignments = {"HF_Team_1": [("Alice", 1), ("Bob", 2)]}
        ko = functions.sync_ko_round(None, "HF", assignments)
        ko.loc[ko["Spieler"] == "Alice", "V1"] = 7

        renamed = {"HF_Team_1": [("Alice Smith", 1), ("Bob", 2)]}
        ko = functions.sync_ko_round(ko, "HF", renamed)
        self.assertEqual(int(ko.loc[ko["PlayerID"] == 1, "V1"].iloc[0]), 7)
        self.assertIn("Alice Smith", set(ko["Spieler"]))

        replaced = {"HF_Team_1": [("Carol", 99), ("Bob", 2)]}
        ko = functions.sync_ko_round(ko, "HF", replaced)
        self.assertEqual(int(ko.loc[ko["PlayerID"] == 99, "V1"].iloc[0]), 0)
        self.assertNotIn("Alice Smith", set(ko["Spieler"]))


class ExcelTests(unittest.TestCase):
    def test_build_workbook_contains_expected_sheets(self):
        df = make_players({"Team A": [("P1", 30, 20)]})
        data = functions.build_workbook(df)
        workbook = load_workbook(io.BytesIO(data), read_only=True)
        self.assertEqual(
            set(workbook.sheetnames),
            {"Alle Würfe", "Team-Ranking", "Spieler-Ranking"},
        )
        self.assertNotIn("_id", [c.value for c in next(workbook["Alle Würfe"].iter_rows())])


class PersistenceIntegrationTests(unittest.TestCase):
    def setUp(self):
        if os.path.exists(functions.DB_PATH):
            os.remove(functions.DB_PATH)

    def test_value_is_saved_and_reloaded(self):
        from streamlit.testing.v1 import AppTest

        app_path = os.path.join(os.path.dirname(__file__), "..", "kegel.py")

        first = AppTest.from_file(app_path, default_timeout=90)
        first.run()
        self.assertEqual(len(first.exception), 0)

        for widget in first.number_input:
            if widget.key == "volle_b1_0_0_1":
                widget.set_value(7)
        first.run()
        self.assertEqual(len(first.exception), 0)

        conn = sqlite3.connect(functions.DB_PATH)
        try:
            count = conn.execute("SELECT COUNT(*) FROM players").fetchone()[0]
            value = conn.execute(
                "SELECT v1 FROM players ORDER BY id LIMIT 1"
            ).fetchone()[0]
        finally:
            conn.close()
        self.assertEqual(count, 50)
        self.assertEqual(value, 7)

        second = AppTest.from_file(app_path, default_timeout=90)
        second.run()
        self.assertEqual(len(second.exception), 0)
        loaded = [
            widget.value
            for widget in second.number_input
            if widget.key == "volle_b1_0_0_1"
        ]
        self.assertEqual(loaded, [7])


if __name__ == "__main__":
    unittest.main()
